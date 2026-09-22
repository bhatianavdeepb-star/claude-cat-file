#!/usr/bin/env python3
"""
Merge a filled-in Five_Plus_Catalogue_Products.xlsx (one tab per catalogue
page) back into index.html: updates existing products' name/sku/spec/price,
adds brand-new products (with a PHOTO_PENDING placeholder image), and creates
new subcategories on the fly when a sheet's Subcategory column names one that
doesn't exist yet.

Usage:
    python3 scripts/apply_catalog_updates.py Five_Plus_Catalogue_Products.xlsx [--html index.html]
"""
import argparse
import re
import sys
import unicodedata

import openpyxl

FIELD_ORDER = ['id', 'category', 'name', 'sku', 'spec', 'price']


def slugify(text):
    text = unicodedata.normalize('NFKD', text).encode('ascii', 'ignore').decode()
    text = re.sub(r"[^a-zA-Z0-9]+", '-', text).strip('-').lower()
    return text or 'item'


def js_escape(text):
    return (text or '').replace('\\', '\\\\').replace("'", "\\'")


def load_html(path):
    with open(path, 'r') as f:
        return f.read()


def find_block(html, start_marker, end_marker):
    start = html.index(start_marker)
    end = html.index(end_marker, start)
    return start, end


def parse_top_categories(html):
    s, e = find_block(html, 'const TOP_CATEGORIES = [', '\n  ];')
    body = html[s + len('const TOP_CATEGORIES = ['):e]
    pages = []
    for m in re.finditer(r"\{\s*id:\s*'([^']+)',\s*label:\s*'([^']*)'.*?sub:\s*\[([^\]]*)\]\s*\}", body):
        pid, label, subs = m.groups()
        sub_list = [x.strip().strip("'") for x in subs.split(',') if x.strip()]
        pages.append({'id': pid, 'label': label, 'sub': sub_list})
    return pages, s, e


def parse_categories(html):
    s, e = find_block(html, 'const CATEGORIES = {', '\n  };\n\n  // Top-level')
    body = html[s + len('const CATEGORIES = {'):e]
    labels = {}
    for m in re.finditer(r"'?([\w-]+)'?:\s*\{\s*\n\s*label:\s*'([^']*)'", body):
        labels[m.group(1)] = m.group(2)
    return labels, s, e


def parse_products(html):
    s = html.index('const PRODUCTS = [')
    e = html.index('\n  ];', s)
    body = html[s + len('const PRODUCTS = ['):e]
    chunks = re.split(r"(?=\{\s*id:)", body)
    products = []
    for chunk in chunks:
        if not chunk.strip():
            continue
        idm = re.search(r"id:\s*'([^']*)'", chunk)
        if not idm:
            continue
        products.append({'id': idm.group(1), 'raw': chunk})
    return products, s, e


def sheet_page_id(sheet_title, pages):
    def norm(x):
        return re.sub(r'[^a-z0-9]', '', unicodedata_lower(x))
    target = norm(sheet_title)
    for p in pages:
        label_plain = p['label'].replace('&amp;', 'and').replace('&', 'and')
        if norm(label_plain) == target:
            return p
    return None


def unicodedata_lower(s):
    return s.lower()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('--html', default='index.html')
    ap.add_argument('--dry-run', action='store_true')
    args = ap.parse_args()

    html = load_html(args.html)
    pages, top_s, top_e = parse_top_categories(html)
    cat_labels, cat_s, cat_e = parse_categories(html)
    products, prod_s, prod_e = parse_products(html)
    existing_ids = {p['id'] for p in products}
    label_to_key = {v.upper(): k for k, v in cat_labels.items()}

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)

    updates = {}          # id -> dict(name, sku, spec, price)
    new_products = []     # list of dicts
    new_categories = {}   # key -> label
    page_new_subs = {}    # page_id -> set(keys) to append to TOP_CATEGORIES sub[]

    for sheet_name in wb.sheetnames:
        if sheet_name == 'READ ME':
            continue
        page = sheet_page_id(sheet_name, pages)
        if not page:
            print(f"WARNING: sheet '{sheet_name}' doesn't match any catalogue page, skipping", file=sys.stderr)
            continue
        ws = wb[sheet_name]
        headers = [c.value for c in ws[1]]
        col = {h: i for i, h in enumerate(headers) if h}

        for row in ws.iter_rows(min_row=2, values_only=True):
            name = row[col['Product Name']] if col.get('Product Name') is not None else None
            if not name or not str(name).strip():
                continue
            name = str(name).strip()
            pid = row[col['Product ID']] if 'Product ID' in col else None
            sku = row[col['SKU']] if 'SKU' in col else None
            spec = row[col['Spec / Size']] if 'Spec / Size' in col else None
            price = row[col['Price (£)']] if 'Price (£)' in col else None
            subcat_label = row[col['Subcategory']] if 'Subcategory' in col else None

            if pid and str(pid).strip():
                pid = str(pid).strip()
                if pid not in existing_ids:
                    print(f"WARNING: unknown Product ID '{pid}' on sheet '{sheet_name}', skipping", file=sys.stderr)
                    continue
                updates[pid] = {
                    'name': name,
                    'sku': '' if sku is None else str(sku).strip(),
                    'spec': '' if spec is None else str(spec).strip(),
                    'price': None if price in (None, '') else float(price),
                }
            else:
                subcat_label = (str(subcat_label).strip() if subcat_label else '') or page['label']
                key = label_to_key.get(subcat_label.upper())
                if not key:
                    key = f"{page['id']}-{slugify(subcat_label)}"
                    if key not in new_categories and key not in cat_labels:
                        new_categories[key] = subcat_label
                    label_to_key[subcat_label.upper()] = key
                if key not in page['sub'] and key not in cat_labels:
                    page_new_subs.setdefault(page['id'], set()).add(key)
                elif key not in page['sub']:
                    page_new_subs.setdefault(page['id'], set()).add(key)

                base_id = f"{page['id']}-{slugify(name)}"
                new_id = base_id
                n = 2
                while new_id in existing_ids or any(p['id'] == new_id for p in new_products):
                    new_id = f"{base_id}-{n}"
                    n += 1

                new_products.append({
                    'id': new_id,
                    'category': key,
                    'name': name,
                    'sku': '' if sku is None else str(sku).strip(),
                    'spec': '' if spec is None else str(spec).strip(),
                    'price': None if price in (None, '') else float(price),
                })

    print(f"Existing products to update: {len(updates)}")
    print(f"New products to add: {len(new_products)}")
    print(f"New subcategories to create: {list(new_categories.keys())}")

    # ---- Apply updates to existing products (edit in place, keep img/pack) ----
    def rewrite_chunk(chunk, upd):
        chunk = re.sub(r"(name:\s*)'((?:\\.|[^'\\])*)'", lambda m: m.group(1) + "'" + js_escape(upd['name']) + "'", chunk, count=1)
        chunk = re.sub(r"(sku:\s*)'((?:\\.|[^'\\])*)'", lambda m: m.group(1) + "'" + js_escape(upd['sku']) + "'", chunk, count=1)
        chunk = re.sub(r"(spec:\s*)'((?:\\.|[^'\\])*)'", lambda m: m.group(1) + "'" + js_escape(upd['spec']) + "'", chunk, count=1)
        price_str = 'null' if upd['price'] is None else repr(upd['price'])
        chunk = re.sub(r"(price:\s*)(null|-?\d+(?:\.\d+)?)", lambda m: m.group(1) + price_str, chunk, count=1)
        return chunk

    new_chunks = []
    for p in products:
        if p['id'] in updates:
            new_chunks.append(rewrite_chunk(p['raw'], updates[p['id']]))
        else:
            new_chunks.append(p['raw'])

    if new_products and new_chunks and not new_chunks[-1].rstrip().endswith(','):
        new_chunks[-1] = new_chunks[-1].rstrip() + ',\n    '

    for np in new_products:
        price_str = 'null' if np['price'] is None else repr(np['price'])
        obj = (f"{{ id:'{np['id']}', category:'{np['category']}', "
               f"name:'{js_escape(np['name'])}', sku:'{js_escape(np['sku'])}', "
               f"spec:'{js_escape(np['spec'])}', price:{price_str}, img:PHOTO_PENDING }},\n    ")
        new_chunks.append(obj)

    products_body = ''.join(new_chunks)
    new_html = html[:prod_s] + 'const PRODUCTS = [' + products_body + html[prod_e:]

    # ---- Add new CATEGORIES entries ----
    if new_categories:
        s, e = find_block(new_html, 'const CATEGORIES = {', '\n  };\n\n  // Top-level')
        insertion = ',\n'
        for key, label in new_categories.items():
            insertion += (f"    '{key}': {{\n      label: '{js_escape(label)}',\n"
                          f"      icon: ICONS.kitchenware,\n      feats: []\n    }},\n")
        insertion = insertion.rstrip('\n')
        new_html = new_html[:e] + insertion + new_html[e:]

    # ---- Update TOP_CATEGORIES sub[] arrays for pages that got new subcats ----
    if page_new_subs:
        s, e = find_block(new_html, 'const TOP_CATEGORIES = [', '\n  ];')
        for page_id, keys in page_new_subs.items():
            m = re.search(r"(\{\s*id:\s*'" + re.escape(page_id) + r"'.*?sub:\s*\[)([^\]]*)(\])", new_html[s:e])
            if not m:
                continue
            existing = [x.strip().strip("'") for x in m.group(2).split(',') if x.strip()]
            for k in keys:
                if k not in existing:
                    existing.append(k)
            new_sub = ','.join(f"'{k}'" for k in existing)
            body = new_html[s:e]
            body = body[:m.start(2)] + new_sub + body[m.end(2):]
            new_html = new_html[:s] + body + new_html[e:]
            s, e = find_block(new_html, 'const TOP_CATEGORIES = [', '\n  ];')

    if args.dry_run:
        print("Dry run only — index.html not written.")
        return

    with open(args.html, 'w') as f:
        f.write(new_html)
    print(f"Wrote {args.html}")


if __name__ == '__main__':
    main()
